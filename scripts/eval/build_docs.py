"""文本 → 文档：MLDR 中文测试集（`Shitao/MLDR` 的 `mldr-v1.0-zh/test.jsonl.gz`）。

- 测试集 800 条查询，每条 1 篇正例和 7 篇难负例（同为维基百科 / 悟道长文）。查询由 LLM 根据正例生成。
- 只取正例不超过 `MAX_CHARS` 字的查询（太长的文档会让评测建索引过慢），抽 80 条；语料为这 80 篇正例
  加每条查询的前 2 篇难负例，去重后约 240 篇。
- 每篇文档按固定顺序轮流写成 `.md`、`.txt`、`.docx`、`.pdf`，用来同时测文档提取链路，标注不变（文档级相关，
  不标页码）。原文没有换行，段落以空格分隔，这里按空格拆段
  （英文单词、编号、数字后的空格不拆，只含符号的碎片并入上一段）；开头空格前的短语是标题，写成文档标题。
- PDF 用 reportlab 嵌入系统黑体（`C:\\Windows\\Fonts\\simhei.ttf`，只用于生成本机评测文件），DOCX 用 python-docx。
  因此只能在 Windows 上构建。
"""

import gzip
import json
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate

from common import hf, reset_dir, rng

REPO = 'Shitao/MLDR'
NAME = 'doc_mldr'
N_QUERIES = 80
N_NEGATIVES = 2
MAX_CHARS = 30000
FONT_PATH = r'C:\Windows\Fonts\simhei.ttf'
FORMATS = ['.md', '.txt', '.docx', '.pdf']


def split_doc(text):
    """拆出标题与段落。"""
    # 前面是 ASCII 字母、数字或“.”的空格不拆：英文词间空格（“New York”）、编号与标题之间
    # （“3.1 元素推断”）、数字与单位之间（“1949 年”）。段落边界前一般是中文标点或汉字，不受影响
    parts = []
    for p in re.split(r'(?<![A-Za-z0-9.]) ', text):
        p = p.strip()
        if not p:
            continue
        # 只有“-”等符号的碎片并入上一段，避免成为单独一行（在 Markdown 中还会被当成列表）
        if parts and not any(c.isalnum() for c in p):
            parts[-1] += p
            continue
        parts.append(p)
    title = parts[0] if parts and len(parts[0]) <= 40 else ''
    return title, parts[1:] if title else parts


def write_md(path, title, paras):
    body = '\n\n'.join(paras)
    path.write_text(f'# {title}\n\n{body}\n' if title else body + '\n', encoding='utf-8')


def write_txt(path, title, paras):
    path.write_text('\n'.join([title, ''] + paras if title else paras) + '\n', encoding='utf-8')


def write_docx(path, title, paras):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if title:
        doc.add_heading(title, level=1)
    for p in paras:
        doc.add_paragraph(p)
    doc.core_properties.title = title
    doc.save(path)


def write_pdf(path, title, paras):
    body = ParagraphStyle('body', fontName='SimHei', fontSize=10.5, leading=16, wordWrap='CJK', spaceAfter=6)
    head = ParagraphStyle('head', parent=body, fontSize=16, leading=24, spaceAfter=12)
    story = [Paragraph(escape(title), head)] if title else []
    story += [Paragraph(escape(p), body) for p in paras]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                            topMargin=2.5 * cm, bottomMargin=2.5 * cm, title=title)
    doc.build(story)


WRITERS = {'.md': write_md, '.txt': write_txt, '.docx': write_docx, '.pdf': write_pdf}


def build(root):
    pdfmetrics.registerFont(TTFont('SimHei', FONT_PATH))
    path = hf(REPO, 'mldr-v1.0-zh/test.jsonl.gz')
    with gzip.open(path, 'rt', encoding='utf-8') as f:
        rows = [json.loads(line) for line in f]
    ok = [r for r in rows if len(r['positive_passages'][0]['text']) <= MAX_CHARS]
    r = rng(NAME)
    picked = sorted(r.sample(ok, N_QUERIES), key=lambda q: int(q['query_id'].rsplit('-', 1)[1]))

    docs = {}
    for q in picked:
        for p in q['positive_passages'] + [n for n in q['negative_passages'] if len(n['text']) <= MAX_CHARS][:N_NEGATIVES]:
            docs[p['docid']] = p['text']

    out = reset_dir(root / 'files' / NAME)
    ids = sorted(docs, key=lambda d: int(d.rsplit('-', 1)[1]))
    order = ids[:]
    r.shuffle(order)
    names = {}
    for i, docid in enumerate(order):
        ext = FORMATS[i % len(FORMATS)]
        names[docid] = f"{docid.removeprefix('doc-zh-').zfill(6)}{ext}"
        title, paras = split_doc(docs[docid])
        WRITERS[ext](out / names[docid], title, paras)

    result = [{
        'id': f"{NAME}-{q['query_id'].removeprefix('q-zh-')}",
        'subset': NAME,
        'query': q['query'].strip(),
        'relevant': [{'file': f"files/{NAME}/{names[p['docid']]}"} for p in q['positive_passages']],
    } for q in picked]

    meta = {
        'name': NAME,
        'modality': 'document',
        'description': '中文长文档（维基百科、悟道语料，平均约 1 万字），轮流存为 Markdown、TXT、DOCX、PDF；查询是针对文档内容的问句，干扰项含 MLDR 提供的难负例。',
        'source': f'https://huggingface.co/datasets/{REPO}（MLDR，BGE-M3 论文）',
        'license': 'MIT（数据集）；文本源自维基百科（CC BY-SA）与悟道语料。只在本机评测。',
        'query_lang': 'zh',
        'query_origin': 'llm',
        'count_files': len(names),
        'count_queries': len(result),
    }
    return meta, result
